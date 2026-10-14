#!/usr/bin/env python3
"""Export layout_dps.xlsx → layout_dps_leaiute.csv + layout_dps.csv (repo root)."""
from __future__ import annotations

import csv
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "layout_dps.xlsx"
LEIAUTE_SHEET = "LEIAUTE DPS_NFS-e "
RN_SHEETS = ("RN DPS_NFS-e", "RN_RECEPCAO_DPS")


def main() -> None:
    if not XLSX.is_file():
        print(f"Missing {XLSX}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

    leiaute_out = ROOT / "layout_dps_leaiute.csv"
    with leiaute_out.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for row in wb[LEIAUTE_SHEET].iter_rows(values_only=True):
            w.writerow(["" if c is None else c for c in row])
    print(f"Wrote {leiaute_out}")

    rn_out = ROOT / "layout_dps.csv"
    with rn_out.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        first_sheet = True
        for sheet in RN_SHEETS:
            ws = wb[sheet]
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                cells = ["" if c is None else c for c in row]
                if i == 0:
                    if first_sheet:
                        w.writerow(["ORIGEM", *cells])
                        first_sheet = False
                    continue
                w.writerow([sheet, *cells])
    print(f"Wrote {rn_out}")
    wb.close()


if __name__ == "__main__":
    main()
